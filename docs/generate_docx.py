import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # University Header
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("CHARUSAT UNIVERSITY\n")
    run_uni.font.name = FONT_NAME
    run_uni.font.size = Pt(16)
    run_uni.font.bold = True
    run_uni.font.color.rgb = COLOR_BLACK

    run_dept = p_uni.add_run("Department of Computer Science & Engineering\n")
    run_dept.font.name = FONT_NAME
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_BLACK

    run_sub = p_uni.add_run("Advance Web Development Frameworks (Semester 5)\n")
    run_sub.font.name = FONT_NAME
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Practical Title Box
    table_title = doc.add_table(rows=1, cols=1)
    table_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_title.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    cell.width = Inches(6.5)
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(6)

    r_pr = p_title.add_run("PRACTICAL REPORT: PRACTICAL NO. 1\n")
    r_pr.font.name = FONT_NAME
    r_pr.font.size = Pt(13)
    r_pr.font.bold = True
    r_pr.font.color.rgb = COLOR_BLACK

    r_pr_name = p_title.add_run("Student Portfolio UI with React & Vite Reusable Components")
    r_pr_name.font.name = FONT_NAME
    r_pr_name.font.size = Pt(12)
    r_pr_name.font.bold = True
    r_pr_name.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Student Name:", "Priyansh Pravinbhai Panchal"),
        ("Degree & Branch:", "B.Tech Computer Science & Engineering (Semester 5)"),
        ("Institute:", "CHARUSAT University"),
        ("Date:", "October 2026")
    ]
    for row_idx, (k, v) in enumerate(info_data):
        c0 = info_table.cell(row_idx, 0)
        c1 = info_table.cell(row_idx, 1)
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F9FAFB")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k)
        r0.font.name = FONT_NAME
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(v)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        return h

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        r_txt = p.add_run(text)
        r_txt.font.name = FONT_NAME
        r_txt.font.size = Pt(11)
        r_txt.font.color.rgb = COLOR_BLACK

    # 1. Objective
    add_section_header("1. Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(4)
    r = p_obj.add_run("To set up a modern React development environment using Vite and build a static UI for a student portfolio page using independently structured, reusable components and props.")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # 2. Prerequisites & Resources
    add_section_header("2. Prerequisites & Resources Required")
    p = doc.add_paragraph()
    add_bullet(p, "• Prerequisites: ", "HTML5, Modern CSS3, JavaScript fundamentals (ES6+ variables, arrow functions, array methods), component-based architecture.")
    p = doc.add_paragraph()
    add_bullet(p, "• Tools & Technologies: ", "Node.js (v24.x), npm (v11.x), Vite 6+, React 18+, Visual Studio Code, Modern Web Browser.")

    # 3. Problem Definition
    add_section_header("3. Problem Definition")
    for req in [
        "1. Create a React application using Vite for a student portfolio page.",
        "2. The application must include at least 4 reusable components: Header, About, Skills, and Footer.",
        "3. Each component must be independently structured and composed into a single-page layout.",
        "4. No logic or JSX should be duplicated across components.",
        "5. Props must be used to pass at least one piece of data into a minimum of 2 components."
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(req)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK

    # 4. Supplementary Problems
    add_section_header("4. Supplementary Problems Solved")
    for sup in [
        ("• Supplementary 1: ", "Added a NavBar component that highlights the active section via state synchronization & scrollspy."),
        ("• Supplementary 2: ", "Passed an array of skills as a prop to the Skills component and rendered dynamically using .map()."),
        ("• Supplementary 3: ", "Added a theme color prop (Royal Blue) to the Header and applied it as an inline style.")
    ]:
        p = doc.add_paragraph()
        add_bullet(p, sup[0], sup[1])

    # 5. Architecture & Hierarchy
    add_section_header("5. Architecture & Component Hierarchy")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.left_indent = Inches(0.25)
    r = p_arch.add_run(
        "App.jsx (Root Layout & State Coordinator)\n"
        "  ├── NavBar.jsx   (Site Navigation & Active Section Highlight)\n"
        "  ├── Header.jsx   (Student Name, Title, Tagline, Inline Theme Style)\n"
        "  ├── About.jsx    (Bio, Education Card, Interests Card)\n"
        "  ├── Skills.jsx   (Dynamic Technical Skills List via Props)\n"
        "  └── Footer.jsx   (Contact Links, GitHub Link, Copyright)"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BLACK

    # 6. Core Implementation Snippets (Concise as requested)
    add_section_header("6. Core Implementation Snippets")

    snippets = [
        (
            "Snippet 1: Dynamic Skills Rendering via Props (Skills.jsx)",
            "Demonstrates dynamic list rendering of skills passed via prop using array.map():",
            """function Skills({ skillList }) {
  return (
    <section id="skills" className="portfolio-section skills-section">
      <h2 className="section-heading">Skills</h2>
      <ul className="skills-list">
        {skillList.map((s) => (
          <li key={s} className="skill-item">
            <span className="skill-dot"></span>
            <span className="skill-name">{s}</span>
          </li>
        ))}
      </ul>
    </section>
  );
}
export default Skills;"""
        ),
        (
            "Snippet 2: Inline Theme Color Prop Application (Header.jsx)",
            "Demonstrates receiving themeColor prop and applying it directly via inline style:",
            """function Header({ name, title, tagline, themeColor = '#2563eb' }) {
  // Supplementary Problem 3: Inline style from themeColor prop
  const headerStyle = {
    backgroundColor: themeColor,
    color: '#ffffff',
  };

  return (
    <header className="portfolio-header" style={headerStyle}>
      <span className="badge">Student Portfolio</span>
      <h1 className="header-name">{name}</h1>
      <h2 className="header-title">{title}</h2>
      {tagline && <p className="header-tagline">{tagline}</p>}
    </header>
  );
}
export default Header;"""
        ),
        (
            "Snippet 3: Active Section Navigation Highlight (NavBar.jsx)",
            "Demonstrates conditional class application based on activeSection prop and aria-current accessibility:",
            """{navItems.map((item) => {
  const isActive = activeSection === item.id;
  return (
    <li key={item.id} className="nav-item">
      <a
        href={`#${item.id}`}
        className={`nav-link ${isActive ? 'active' : ''}`}
        aria-current={isActive ? 'page' : undefined}
        onClick={(e) => handleNavClick(e, item.id)}
      >
        {item.label}
      </a>
    </li>
  );
})}"""
        ),
        (
            "Snippet 4: Component Composition & Props Passing (App.jsx)",
            "Demonstrates composing independent reusable components into a single layout and passing data down via props:",
            """return (
  <div className="portfolio-app">
    <NavBar brand={student.name} activeSection={activeSection} onSelectSection={setActiveSection} />
    <Header name={student.name} title={student.title} tagline={student.tagline} themeColor={student.themeColor} />
    <main className="portfolio-content">
      <About bio={student.bio} education={student.education} interests={student.interests} />
      <Skills skillList={student.skills} />
    </main>
    <Footer name={student.name} year={2026} email={student.email} github={student.github} />
  </div>
);"""
        ),
    ]

    for title, desc, code in snippets:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(2)
        r_f = p_title.add_run(title)
        r_f.font.name = FONT_NAME
        r_f.font.size = Pt(11)
        r_f.font.bold = True
        r_f.font.color.rgb = COLOR_BLACK

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(3)
        r_d = p_desc.add_run(desc)
        r_d.font.name = FONT_NAME
        r_d.font.size = Pt(10)
        r_d.font.italic = True
        r_d.font.color.rgb = COLOR_BLACK

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        set_cell_background(c, "F8FAFC")
        c.width = Inches(6.5)
        set_cell_margins(c, top=60, bottom=60, left=100, right=100)

        p_code = c.paragraphs[0]
        p_code.paragraph_format.space_before = Pt(2)
        p_code.paragraph_format.space_after = Pt(2)
        r_c = p_code.add_run(code)
        r_c.font.name = "Consolas"
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = COLOR_BLACK

    # 7. Output & Verification
    add_section_header("7. Output & Verification")
    p_out = doc.add_paragraph()
    r = p_out.add_run(
        "• Linter (oxlint): Passed with 0 errors and 0 warnings.\n"
        "• Vite Build: Compiled successfully in 110ms with zero errors.\n"
        "• Features Verified: Dynamic skill badges mapped cleanly, sticky active NavBar highlighting current section, inline Royal Blue theme applied, responsive cards, and live contact links."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # Screenshot placeholder
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(8)
    p_box.paragraph_format.space_after = Pt(8)
    
    tbl_shot = doc.add_table(rows=1, cols=1)
    tbl_shot.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_shot = tbl_shot.cell(0, 0)
    set_cell_background(c_shot, "F3F4F6")
    c_shot.width = Inches(6.5)
    set_cell_margins(c_shot, top=180, bottom=180, left=150, right=150)
    p_sh = c_shot.paragraphs[0]
    p_sh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sh = p_sh.add_run("[ INSERT BROWSER SCREENSHOT OF PORTFOLIO APPLICATION HERE ]")
    r_sh.font.name = FONT_NAME
    r_sh.font.size = Pt(11)
    r_sh.font.bold = True
    r_sh.font.color.rgb = COLOR_BLACK

    # 8. Conclusion
    add_section_header("8. Conclusion")
    p_con = doc.add_paragraph()
    r = p_con.add_run(
        "In this practical, a Vite React development environment was successfully initialized. A modular, responsive student portfolio UI was constructed with independently structured, reusable components. Props were utilized to pass personal data, dynamic skill arrays, and inline theme styles, satisfying all primary and supplementary requirements."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    output_path = r"D:\SEM-5\Advance Web Development Frameworks\docs\Practical-01-Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated Practical 1 Word document at: {output_path}")

if __name__ == "__main__":
    create_document()
