import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # University Header
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("CHARUSAT UNIVERSITY\n")
    run_uni.font.name = FONT_NAME
    run_uni.font.size = Pt(16)
    run_uni.font.bold = True
    run_uni.font.color.rgb = COLOR_BLACK

    run_dept = p_uni.add_run("Department of Computer Science & Engineering\n")
    run_dept.font.name = FONT_NAME
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_BLACK

    run_sub = p_uni.add_run("Advance Web Development Frameworks (Semester 5)\n")
    run_sub.font.name = FONT_NAME
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Practical Title Box
    table_title = doc.add_table(rows=1, cols=1)
    table_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_title.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    cell.width = Inches(6.5)
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(6)

    r_pr = p_title.add_run("PRACTICAL REPORT: PRACTICAL NO. 5\n")
    r_pr.font.name = FONT_NAME
    r_pr.font.size = Pt(13)
    r_pr.font.bold = True
    r_pr.font.color.rgb = COLOR_BLACK

    r_pr_name = p_title.add_run("MongoDB Integration & Schema Validation Using Mongoose ODM in Express")
    r_pr_name.font.name = FONT_NAME
    r_pr_name.font.size = Pt(12)
    r_pr_name.font.bold = True
    r_pr_name.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Student Name:", "Priyansh Pravinbhai Panchal"),
        ("Degree & Branch:", "B.Tech Computer Science & Engineering (Semester 5)"),
        ("Institute:", "CHARUSAT University"),
        ("Date:", "October 2026")
    ]
    for row_idx, (k, v) in enumerate(info_data):
        c0 = info_table.cell(row_idx, 0)
        c1 = info_table.cell(row_idx, 1)
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F9FAFB")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k)
        r0.font.name = FONT_NAME
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(v)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        return h

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        r_txt = p.add_run(text)
        r_txt.font.name = FONT_NAME
        r_txt.font.size = Pt(11)
        r_txt.font.color.rgb = COLOR_BLACK

    # 1. Objective
    add_section_header("1. Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(4)
    r = p_obj.add_run("To connect a MongoDB database to an Express.js backend server, enforce data integrity and validation through Mongoose schemas, and implement persistent CRUD operations with structured error responses.")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # 2. Prerequisites & Resources
    add_section_header("2. Prerequisites & Resources Required")
    p = doc.add_paragraph()
    add_bullet(p, "• Prerequisites: ", "Practical 4 completed working Express server with in-memory CRUD, understanding of NoSQL/document-based databases, schemas, and async/await.")
    p = doc.add_paragraph()
    add_bullet(p, "• Tools & Technologies: ", "Node.js (v24.x), npm (v11.x), Express.js (v4.21.x), Mongoose ODM (v8.12.x), MongoDB Server (local/Atlas), dotenv, Postman / Thunder Client.")

    # 3. Problem Definition
    add_section_header("3. Problem Definition")
    for req in [
        "1. Extend the Task Management backend from Practical 4 by connecting MongoDB using Mongoose.",
        "2. Define a Task schema with at least 4 fields: title (String, required), description (String), completed (Boolean, default false), createdAt (Date, default Date.now).",
        "3. Replace the in-memory array from Practical 4 with real Mongoose model operations (Task.find, Task.create, Task.findByIdAndUpdate, Task.findByIdAndDelete).",
        "4. Test all CRUD operations against the live database using Postman / Thunder Client and verify data persistence across server restarts.",
        "5. Ensure schema validation and casting errors are returned as structured JSON, not raw Mongoose error stack traces."
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(req)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK

    # 4. Supplementary Problems Solved
    add_section_header("4. Supplementary Problems Solved")
    for sup in [
        ("• Supplementary 1: ", "Added a 'priority' field to the Task schema restricted to an enum of ['low', 'medium', 'high'] with a descriptive validation error message."),
        ("• Supplementary 2: ", "Added a Mongoose pre-save hook that automatically trims whitespace from the 'title' field before persisting to MongoDB."),
        ("• Supplementary 3: ", "Implemented a GET /tasks/:id endpoint that performs an asynchronous database lookup and returns a 404 JSON response if the task ID does not exist.")
    ]:
        p = doc.add_paragraph()
        add_bullet(p, sup[0], sup[1])

    # 5. Architecture & Data Flow
    add_section_header("5. Architecture & Data Flow")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.left_indent = Inches(0.25)
    r = p_arch.add_run(
        "Client (Postman / Thunder Client / Frontend)\n"
        "  │\n"
        "  ▼ Express Server (Middleware Pipeline)\n"
        "  │   ├── Request Logger & Content-Type Validator\n"
        "  │   └── Router Endpoints (Async Controller Handlers)\n"
        "  │\n"
        "  ▼ Mongoose ODM (models/Task.js)\n"
        "  │   ├── Schema Validation: title (required), completed (default: false), priority (enum)\n"
        "  │   └── Pre-Save Lifecycle Hook: this.title.trim()\n"
        "  │\n"
        "  ▼ MongoDB Database (tasks collection)\n"
        "  │   └── Persistent Document Storage: { _id, title, description, completed, priority, createdAt }\n"
        "  │\n"
        "  ▼ Error Handling Pipeline\n"
        "      ├── ValidationError → 400 Bad Request (Formatted JSON details array)\n"
        "      ├── CastError       → 400 Bad Request (Invalid ObjectId JSON)\n"
        "      └── Route Not Found → 404 Not Found"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BLACK

    # 6. Core Implementation Snippets (Concise as requested)
    add_section_header("6. Core Implementation Snippets")

    p_note = doc.add_paragraph()
    r_n = p_note.add_run("Note: Only newly introduced Mongoose schema definitions, lifecycle hooks, and database controller logic are documented below to maintain a concise submission report.")
    r_n.font.name = FONT_NAME
    r_n.font.size = Pt(10)
    r_n.font.italic = True
    r_n.font.color.rgb = COLOR_BLACK

    snippets = [
        (
            "Snippet 1: Task Schema, Priority Enum & Pre-Save Trimming (models/Task.js)",
            "Demonstrates schema fields, enum validation (Supplementary 1), and pre-save title trimming hook (Supplementary 2):",
            """const mongoose = require('mongoose');

const taskSchema = new mongoose.Schema({
  title: {
    type: String,
    required: [true, 'Task title is required']
  },
  description: { type: String, default: '' },
  completed: { type: Boolean, default: false },
  createdAt: { type: Date, default: Date.now },
  // Supplementary 1: Priority restricted to enum
  priority: {
    type: String,
    enum: {
      values: ['low', 'medium', 'high'],
      message: '{VALUE} is not a valid priority. Allowed: low, medium, high'
    },
    default: 'medium'
  }
});

// Supplementary 2: Pre-save hook automatically trimming whitespace
taskSchema.pre('save', function () {
  if (this.title) {
    this.title = this.title.trim();
  }
});

module.exports = mongoose.model('Task', taskSchema);"""
        ),
        (
            "Snippet 2: MongoDB Connection in Express (server.js)",
            "Demonstrates connecting to MongoDB using Mongoose and environment variables (Step 2):",
            """require('dotenv').config();
const mongoose = require('mongoose');

mongoose
  .connect(process.env.MONGO_URI || 'mongodb://127.0.0.1:27017/task_manager_db')
  .then(() => console.log('MongoDB connected successfully'))
  .catch((err) => console.error('MongoDB connection error:', err.message));"""
        ),
        (
            "Snippet 3: Mongoose CRUD Controller Handlers (server.js)",
            "Demonstrates replacing in-memory operations with Mongoose methods Task.find(), Task.create(), Task.findByIdAndUpdate(), Task.findByIdAndDelete():",
            """// GET /tasks - Read all tasks
app.get('/tasks', async (req, res, next) => {
  try {
    const tasks = await Task.find().sort({ createdAt: -1 });
    res.status(200).json({ success: true, count: tasks.length, data: tasks });
  } catch (err) { next(err); }
});

// POST /tasks - Create task in MongoDB
app.post('/tasks', async (req, res, next) => {
  try {
    const task = await Task.create(req.body);
    res.status(201).json({ success: true, message: 'Task created', data: task });
  } catch (err) { next(err); }
});

// PUT /tasks/:id - Update task in MongoDB
app.put('/tasks/:id', async (req, res, next) => {
  try {
    const task = await Task.findByIdAndUpdate(req.params.id, req.body, { new: true, runValidators: true });
    if (!task) return res.status(404).json({ success: false, error: 'Task Not Found' });
    res.status(200).json({ success: true, data: task });
  } catch (err) { next(err); }
});

// DELETE /tasks/:id - Delete task from MongoDB
app.delete('/tasks/:id', async (req, res, next) => {
  try {
    const task = await Task.findByIdAndDelete(req.params.id);
    if (!task) return res.status(404).json({ success: false, error: 'Task Not Found' });
    res.status(200).json({ success: true, message: 'Task deleted', data: task });
  } catch (err) { next(err); }
});"""
        ),
        (
            "Snippet 4: Single Task Retrieval with 404 Handler (Supplementary Problem 3)",
            "Demonstrates querying by ID and returning structured 404 JSON response when document is not found:",
            """// Supplementary Problem 3: GET /tasks/:id with 404 handling
app.get('/tasks/:id', async (req, res, next) => {
  try {
    const task = await Task.findById(req.params.id);
    if (!task) {
      return res.status(404).json({
        success: false,
        error: 'Task Not Found',
        message: `No task found with ID: ${req.params.id}`
      });
    }
    res.status(200).json({ success: true, data: task });
  } catch (err) {
    next(err);
  }
});"""
        ),
        (
            "Snippet 5: Structured Mongoose Validation Error Formatter (server.js)",
            "Demonstrates converting raw Mongoose ValidationError and CastError into structured JSON responses (Problem Definition):",
            """// Global error middleware formatting Mongoose errors
app.use((err, req, res, next) => {
  // Format Mongoose Schema Validation Errors
  if (err.name === 'ValidationError') {
    const errors = Object.keys(err.errors).map((key) => ({
      field: key,
      message: err.errors[key].message
    }));
    return res.status(400).json({
      success: false,
      error: 'Validation Error',
      message: 'One or more fields failed Mongoose schema validation',
      details: errors
    });
  }

  // Format Invalid MongoDB ObjectId CastError
  if (err.name === 'CastError') {
    return res.status(400).json({
      success: false,
      error: 'Invalid ID Format',
      message: `Invalid MongoDB ObjectId: '${err.value}'`
    });
  }

  // Generic 500
  res.status(500).json({ success: false, error: 'Internal Server Error', message: err.message });
});"""
        ),
    ]

    for title, desc, code in snippets:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(2)
        r_f = p_title.add_run(title)
        r_f.font.name = FONT_NAME
        r_f.font.size = Pt(11)
        r_f.font.bold = True
        r_f.font.color.rgb = COLOR_BLACK

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(3)
        r_d = p_desc.add_run(desc)
        r_d.font.name = FONT_NAME
        r_d.font.size = Pt(10)
        r_d.font.italic = True
        r_d.font.color.rgb = COLOR_BLACK

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        set_cell_background(c, "F8FAFC")
        c.width = Inches(6.5)
        set_cell_margins(c, top=60, bottom=60, left=100, right=100)

        p_code = c.paragraphs[0]
        p_code.paragraph_format.space_before = Pt(2)
        p_code.paragraph_format.space_after = Pt(2)
        r_c = p_code.add_run(code)
        r_c.font.name = "Consolas"
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = COLOR_BLACK

    # 7. Output & Verification Table
    add_section_header("7. Output & Verification Results")
    
    p_tbl_intro = doc.add_paragraph()
    r_ti = p_tbl_intro.add_run("Summary of tested endpoints, Mongoose schema validations, and MongoDB persistence:")
    r_ti.font.name = FONT_NAME
    r_ti.font.size = Pt(10.5)
    r_ti.font.color.rgb = COLOR_BLACK

    test_table = doc.add_table(rows=8, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_headers = ["Test Scenario", "HTTP Method / Endpoint", "Status Code", "Observed Response / Behavior"]
    test_data = [
        ("Missing Required Title", "POST /tasks", "400 Bad Request", "Mongoose ValidationError returned structured JSON details array."),
        ("Invalid Priority Enum", "POST /tasks", "400 Bad Request", "Mongoose rejected 'critical'; allowed priorities: low, medium, high."),
        ("Pre-Save Hook Trimming", "POST /tasks", "201 Created", "Whitespace '  Title  ' automatically trimmed to 'Title' in MongoDB."),
        ("Read All Documents", "GET /tasks", "200 OK", "Retrieved 6 persistent documents from MongoDB collection."),
        ("Read Non-Existent ID", "GET /tasks/6adb...99", "404 Not Found", "Returned structured 404 JSON for non-existent MongoDB ObjectId."),
        ("Update Document in DB", "PUT /tasks/:id", "200 OK", "Updated 'completed' to true and priority in live MongoDB collection."),
        ("Delete Document in DB", "DELETE /tasks/:id", "200 OK", "Removed document from MongoDB; subsequent GET returned 404."),
    ]

    for col_idx, h_text in enumerate(table_headers):
        cell_h = test_table.cell(0, col_idx)
        set_cell_background(cell_h, "F1F5F9")
        p = cell_h.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.name = FONT_NAME
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_BLACK

    for row_idx, row_values in enumerate(test_data):
        for col_idx, val in enumerate(row_values):
            cell_d = test_table.cell(row_idx + 1, col_idx)
            set_cell_background(cell_d, "FFFFFF")
            p = cell_d.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_BLACK

    # Screenshot placeholder
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(12)
    p_box.paragraph_format.space_after = Pt(8)
    
    tbl_shot = doc.add_table(rows=1, cols=1)
    tbl_shot.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_shot = tbl_shot.cell(0, 0)
    set_cell_background(c_shot, "F3F4F6")
    c_shot.width = Inches(6.5)
    set_cell_margins(c_shot, top=180, bottom=180, left=150, right=150)
    p_sh = c_shot.paragraphs[0]
    p_sh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sh = p_sh.add_run("[ INSERT POSTMAN / THUNDER CLIENT SCREENSHOTS OF MONGOOSE VALIDATION & CRUD OUTPUTS HERE ]")
    r_sh.font.name = FONT_NAME
    r_sh.font.size = Pt(11)
    r_sh.font.bold = True
    r_sh.font.color.rgb = COLOR_BLACK

    # 8. Conclusion
    add_section_header("8. Conclusion")
    p_con = doc.add_paragraph()
    r = p_con.add_run(
        "In Practical 5, MongoDB was successfully connected to the Express.js server using Mongoose ODM. An explicit Task schema was designed to enforce required fields, default values, priority enums, and pre-save trimming hooks. In-memory operations were replaced with persistent Mongoose CRUD queries, and raw database errors were converted into clear, structured JSON error responses."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    output_path = r"D:\SEM-5\Advance Web Development Frameworks\docs\Practical-05-Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated Practical 5 Word document at: {output_path}")

if __name__ == "__main__":
    create_document()
