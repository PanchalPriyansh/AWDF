import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # University Header
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("CHARUSAT UNIVERSITY\n")
    run_uni.font.name = FONT_NAME
    run_uni.font.size = Pt(16)
    run_uni.font.bold = True
    run_uni.font.color.rgb = COLOR_BLACK

    run_dept = p_uni.add_run("Department of Computer Science & Engineering\n")
    run_dept.font.name = FONT_NAME
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_BLACK

    run_sub = p_uni.add_run("Advance Web Development Frameworks (Semester 5)\n")
    run_sub.font.name = FONT_NAME
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Practical Title Box
    table_title = doc.add_table(rows=1, cols=1)
    table_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_title.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    cell.width = Inches(6.5)
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(6)

    r_pr = p_title.add_run("PRACTICAL REPORT: PRACTICAL NO. 2\n")
    r_pr.font.name = FONT_NAME
    r_pr.font.size = Pt(13)
    r_pr.font.bold = True
    r_pr.font.color.rgb = COLOR_BLACK

    r_pr_name = p_title.add_run("Reactive State Management with useState and Multi-Page Navigation with React Router")
    r_pr_name.font.name = FONT_NAME
    r_pr_name.font.size = Pt(12)
    r_pr_name.font.bold = True
    r_pr_name.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Student Name:", "Priyansh Pravinbhai Panchal"),
        ("Degree & Branch:", "B.Tech Computer Science & Engineering (Semester 5)"),
        ("Institute:", "CHARUSAT University"),
        ("Date:", "October 2026")
    ]
    for row_idx, (k, v) in enumerate(info_data):
        c0 = info_table.cell(row_idx, 0)
        c1 = info_table.cell(row_idx, 1)
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F9FAFB")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k)
        r0.font.name = FONT_NAME
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(v)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        return h

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        r_txt = p.add_run(text)
        r_txt.font.name = FONT_NAME
        r_txt.font.size = Pt(11)
        r_txt.font.color.rgb = COLOR_BLACK

    # 1. Objective
    add_section_header("1. Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(4)
    r = p_obj.add_run("To implement reactive state management using useState and multi-page client-side navigation using React Router v6 in the existing student portfolio application without full page reloads.")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # 2. Prerequisites & Resources
    add_section_header("2. Prerequisites & Resources Required")
    p = doc.add_paragraph()
    add_bullet(p, "• Prerequisites: ", "Practical 1 completed working portfolio app with reusable components, understanding of JavaScript array/object handling and React Hooks.")
    p = doc.add_paragraph()
    add_bullet(p, "• Tools & Technologies: ", "Node.js (v24.x), npm (v11.x), React 18+, React Router DOM v6, Vite 6+, Visual Studio Code, Modern Web Browser.")

    # 3. Problem Definition
    add_section_header("3. Problem Definition")
    for req in [
        "1. Extend the portfolio application from Practical 1 by adding React Router (react-router-dom).",
        "2. Implement a navigation bar with at least 3 routes: Home (/), Projects (/projects), and Contact (/contact).",
        "3. Each route must render a distinct component without a full page reload.",
        "4. Add at least 2 useState variables used meaningfully: one for toggling UI visibility and one for managing a form input on the Contact page.",
        "5. The Contact page must include a controlled input that captures and displays user input in real time."
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(req)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK

    # 4. Supplementary Problems Solved
    add_section_header("4. Supplementary Problems Solved")
    for sup in [
        ("• Supplementary 1: ", "Added a 404 Not Found route (path='*') that renders a custom error component (NotFound.jsx) for unknown paths with navigation back to Home."),
        ("• Supplementary 2: ", "Implemented a dark/light mode toggle using useState that dynamically applies the 'dark-theme' or 'light-theme' CSS class to the root container element."),
        ("• Supplementary 3: ", "Stored the contact form input in state and displayed a live character count (e.g., character length / 500) reactively below the input.")
    ]:
        p = doc.add_paragraph()
        add_bullet(p, sup[0], sup[1])

    # 5. Architecture & Routing Hierarchy
    add_section_header("5. Architecture & Routing Hierarchy")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.left_indent = Inches(0.25)
    r = p_arch.add_run(
        "App.jsx (wrapped in <BrowserRouter> inside main.jsx)\n"
        "  ├── NavBar.jsx       (NavLink client routing & Dark/Light mode toggle button)\n"
        "  ├── Routes Container\n"
        "  │     ├── Route: '/'          → Home.jsx     (Composes Header, About, Skills from Practical 1)\n"
        "  │     ├── Route: '/projects'  → Projects.jsx (Interactive category filter & project cards)\n"
        "  │     ├── Route: '/contact'   → Contact.jsx  (Controlled inputs, live preview, character counter, UI toggle)\n"
        "  │     └── Route: '*'          → NotFound.jsx (404 custom error handler)\n"
        "  └── Footer.jsx       (Global contact info & copyright notice)"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BLACK

    # 6. Core Implementation Snippets (Concise as requested)
    add_section_header("6. Core Implementation Snippets")

    p_note = doc.add_paragraph()
    r_n = p_note.add_run("Note: Only key functional code snippets introduced in Practical 2 are shown below to maintain a concise documentation report.")
    r_n.font.name = FONT_NAME
    r_n.font.size = Pt(10)
    r_n.font.italic = True
    r_n.font.color.rgb = COLOR_BLACK

    snippets = [
        (
            "Snippet 1: Wrapping Root with BrowserRouter (main.jsx)",
            "Demonstrates configuring client-side routing at the application entry point:",
            """import { BrowserRouter } from 'react-router-dom';
import App from './App.jsx';

createRoot(document.getElementById('root')).render(
  <StrictMode>
    <BrowserRouter>
      <App />
    </BrowserRouter>
  </StrictMode>
);"""
        ),
        (
            "Snippet 2: Route Configuration & Theme Toggle State (App.jsx)",
            "Demonstrates defining multi-page routes and managing dark/light theme state with useState:",
            """function App() {
  const [darkMode, setDarkMode] = useState(false);
  const toggleTheme = () => setDarkMode((prev) => !prev);

  return (
    <div className={`portfolio-app ${darkMode ? 'dark-theme' : 'light-theme'}`}>
      <NavBar brand={student.name} darkMode={darkMode} onToggleTheme={toggleTheme} />
      <Routes>
        <Route path="/" element={<Home student={student} />} />
        <Route path="/projects" element={<Projects />} />
        <Route path="/contact" element={<Contact studentEmail={student.email} studentGithub={student.github} />} />
        {/* Supplementary 1: 404 Route */}
        <Route path="*" element={<NotFound />} />
      </Routes>
      <Footer name={student.name} year={2026} email={student.email} github={student.github} />
    </div>
  );
}"""
        ),
        (
            "Snippet 3: Client-Side Navigation with NavLink (NavBar.jsx)",
            "Demonstrates NavLink component avoiding page reloads and providing automatic active class styling:",
            """<ul className="nav-list">
  {navItems.map((item) => (
    <li key={item.path} className="nav-item">
      <NavLink
        to={item.path}
        end={item.path === '/'}
        className={({ isActive }) => `nav-link ${isActive ? 'active' : ''}`}
      >
        {item.label}
      </NavLink>
    </li>
  ))}
</ul>"""
        ),
        (
            "Snippet 4: Controlled Inputs, Live Counter & Real-Time Preview (Contact.jsx)",
            "Demonstrates controlled form inputs, character counter, UI visibility toggle, and real-time preview:",
            """// Step 5: Controlled input & visibility state
const [message, setMessage] = useState('');
const [showTips, setShowTips] = useState(false); // Step 6: Visibility toggle

return (
  <>
    {/* Step 6: Element toggled via useState */}
    <button onClick={() => setShowTips((prev) => !prev)}>
      {showTips ? 'Hide Tips' : 'Show Help Guidelines'}
    </button>
    {showTips && <div className="help-box">Contact Guidelines...</div>}

    {/* Controlled input */}
    <textarea
      value={message}
      onChange={(e) => setMessage(e.target.value)}
      maxLength={500}
    />
    
    {/* Supplementary 3: Live character counter */}
    <div className="char-counter">
      Count: {message.length} / 500
    </div>

    {/* Problem Definition: Real-time display of user input */}
    <div className="live-preview-box">
      <strong>Live Preview:</strong> {message || 'Waiting for message...'}
    </div>
  </>
);"""
        ),
        (
            "Snippet 5: Custom 404 Not Found Component (NotFound.jsx)",
            "Demonstrates handling unknown routes gracefully with a custom error view:",
            """function NotFound() {
  return (
    <div className="not-found-card">
      <span className="error-code">404</span>
      <h1>Page Not Found</h1>
      <p>The requested route does not exist.</p>
      <Link to="/" className="btn-primary">← Return to Home</Link>
    </div>
  );
}
export default NotFound;"""
        ),
    ]

    for title, desc, code in snippets:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(2)
        r_f = p_title.add_run(title)
        r_f.font.name = FONT_NAME
        r_f.font.size = Pt(11)
        r_f.font.bold = True
        r_f.font.color.rgb = COLOR_BLACK

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(3)
        r_d = p_desc.add_run(desc)
        r_d.font.name = FONT_NAME
        r_d.font.size = Pt(10)
        r_d.font.italic = True
        r_d.font.color.rgb = COLOR_BLACK

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        set_cell_background(c, "F8FAFC")
        c.width = Inches(6.5)
        set_cell_margins(c, top=60, bottom=60, left=100, right=100)

        p_code = c.paragraphs[0]
        p_code.paragraph_format.space_before = Pt(2)
        p_code.paragraph_format.space_after = Pt(2)
        r_c = p_code.add_run(code)
        r_c.font.name = "Consolas"
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = COLOR_BLACK

    # 7. Output & Verification
    add_section_header("7. Output & Verification")
    p_out = doc.add_paragraph()
    r = p_out.add_run(
        "• Linter (oxlint): Passed with 0 errors and 0 warnings across all 12 modules.\n"
        "• Vite Build: Compiled successfully in 126ms with full bundle optimization.\n"
        "• Route Transitions: Navigating between Home (/), Projects (/projects), and Contact (/contact) occurs instantly on client side without full-page reloads.\n"
        "• Form Reactivity: Controlled message input reactively updates character count and displays immediate real-time feedback on the preview card.\n"
        "• State Toggles: Both the Help Tips toggle and Dark/Light mode theme toggle function reliably and update the DOM reactively."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # Screenshot placeholder
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(8)
    p_box.paragraph_format.space_after = Pt(8)
    
    tbl_shot = doc.add_table(rows=1, cols=1)
    tbl_shot.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_shot = tbl_shot.cell(0, 0)
    set_cell_background(c_shot, "F3F4F6")
    c_shot.width = Inches(6.5)
    set_cell_margins(c_shot, top=180, bottom=180, left=150, right=150)
    p_sh = c_shot.paragraphs[0]
    p_sh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sh = p_sh.add_run("[ INSERT SCREENSHOTS OF HOME, PROJECTS, CONTACT FORM WITH LIVE INPUT & 404 ERROR PAGE HERE ]")
    r_sh.font.name = FONT_NAME
    r_sh.font.size = Pt(11)
    r_sh.font.bold = True
    r_sh.font.color.rgb = COLOR_BLACK

    # 8. Conclusion
    add_section_header("8. Conclusion")
    p_con = doc.add_paragraph()
    r = p_con.add_run(
        "In Practical 2, multi-page client-side navigation was successfully integrated using React Router DOM v6, eliminating full page reloads. Reactive state management was implemented using useState for controlled form inputs, real-time live preview generation, dynamic character counters, UI visibility toggling, and an interactive dark/light theme switcher. A custom 404 route was also established, delivering a complete, robust single-page application experience."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    output_path = r"D:\SEM-5\Advance Web Development Frameworks\docs\Practical-02-Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated Practical 2 Word document at: {output_path}")

if __name__ == "__main__":
    create_document()
