import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # University Header
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("CHARUSAT UNIVERSITY\n")
    run_uni.font.name = FONT_NAME
    run_uni.font.size = Pt(16)
    run_uni.font.bold = True
    run_uni.font.color.rgb = COLOR_BLACK

    run_dept = p_uni.add_run("Department of Computer Science & Engineering\n")
    run_dept.font.name = FONT_NAME
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_BLACK

    run_sub = p_uni.add_run("Advance Web Development Frameworks (Semester 5)\n")
    run_sub.font.name = FONT_NAME
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Practical Title Box
    table_title = doc.add_table(rows=1, cols=1)
    table_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_title.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    cell.width = Inches(6.5)
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(6)

    r_pr = p_title.add_run("PRACTICAL REPORT: PRACTICAL NO. 8\n")
    r_pr.font.name = FONT_NAME
    r_pr.font.size = Pt(13)
    r_pr.font.bold = True
    r_pr.font.color.rgb = COLOR_BLACK

    r_pr_name = p_title.add_run("Frontend Performance Optimization: Lazy Loading and Code Splitting")
    r_pr_name.font.name = FONT_NAME
    r_pr_name.font.size = Pt(12)
    r_pr_name.font.bold = True
    r_pr_name.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Student Name:", "Priyansh Pravinbhai Panchal"),
        ("Degree & Branch:", "B.Tech Computer Science & Engineering (Semester 5)"),
        ("Institute:", "CHARUSAT University"),
        ("Date:", "October 2026")
    ]
    for row_idx, (k, v) in enumerate(info_data):
        c0 = info_table.cell(row_idx, 0)
        c1 = info_table.cell(row_idx, 1)
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F9FAFB")
        set_cell_background(c1, "FFFFFF")

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k)
        r0.font.name = FONT_NAME
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(v)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        return h

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        r_txt = p.add_run(text)
        r_txt.font.name = FONT_NAME
        r_txt.font.size = Pt(11)
        r_txt.font.color.rgb = COLOR_BLACK

    # 1. Objective
    add_section_header("1. Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(4)
    r = p_obj.add_run("To optimize React frontend performance by implementing route-based code splitting with React.lazy() and Suspense, dynamically loading heavy third-party libraries on demand, preventing loading flicker with minimum-delay fallbacks, and eliminating unnecessary re-renders using the React DevTools Profiler.")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # 2. Prerequisites & Resources
    add_section_header("2. Prerequisites & Resources Required")
    p = doc.add_paragraph()
    add_bullet(p, "• Prerequisites: ", "Practicals 1–7 completed full-stack React + Express + MongoDB application, understanding of browser script evaluation, ES modules dynamic import(), and bundle chunking.")
    p = doc.add_paragraph()
    add_bullet(p, "• Tools & Technologies: ", "Node.js (v24.x), npm, React 18+, Vite (v8.x), Chart.js (v4.x), Chrome DevTools (Network & Performance / Profiler tabs).")

    # 3. Problem Definition
    add_section_header("3. Problem Definition")
    for req in [
        "1. Optimize the React frontend of the Task Management application by implementing lazy loading for route-based components.",
        "2. Apply React.lazy() and Suspense to the Projects, Contact, Home, and NotFound routes.",
        "3. Provide a meaningful fallback UI while a lazy-loaded chunk is loading over the network.",
        "4. Measure and compare bundle size and load time before and after optimization using Vite build output and DevTools.",
        "5. Document the before/after performance comparison with recorded values and metrics."
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(req)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK

    # 4. Supplementary Problems Solved
    add_section_header("4. Supplementary Problems Solved")
    for sup in [
        ("• Supplementary 1: ", "Lazy-loaded a heavy third-party component (Chart.js task analytics, 205.77 KB) in an isolated chunk (TaskAnalyticsChart.jsx) that is fetched strictly on demand when toggled by the user."),
        ("• Supplementary 2: ", "Implemented a custom 300ms minimum-delay fallback component (DelayedFallback.jsx) to eliminate visual flickering on fast network connections where chunks resolve in under 300ms."),
        ("• Supplementary 3: ", "Used the React DevTools Profiler to identify unnecessary re-renders of the Footer and NavBar during parent theme state toggles, and resolved them by wrapping components in React.memo and memoizing callbacks with useCallback.")
    ]:
        p = doc.add_paragraph()
        add_bullet(p, sup[0], sup[1])

    # 5. Architecture & Code-Splitting Bundle Flow
    add_section_header("5. Architecture & Code-Splitting Bundle Flow")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.left_indent = Inches(0.25)
    r = p_arch.add_run(
        "Before Optimization (Monolithic Single Bundle):\n"
        "  main.bundle.js (289.91 KB) ────────► [Home][Projects][Contact][NotFound] all loaded upfront\n\n"
        "After Optimization (Route-Based Code Splitting & Dynamic Imports):\n"
        "  index-CcahhOss.js (263.68 KB)       ──► Core shell loaded on initial page visit\n"
        "    ├── Home-Co8PPuXq.js (2.71 KB)    ──► Loaded ONLY when navigating to '/'\n"
        "    ├── Contact-8lET-bAE.js (4.99 KB) ──► Loaded ONLY when navigating to '/contact'\n"
        "    ├── Projects-B8jro97-.js (21.49 KB)─► Loaded ONLY when navigating to '/projects'\n"
        "    ├── NotFound-DX-wgjGn.js (0.71 KB)──► Loaded ONLY on 404 routes\n"
        "    └── TaskAnalyticsChart-DE_jqqtA.js (205.77 KB) ──► Chart.js chunk loaded ONLY on demand\n\n"
        "Network Fallback Optimization:\n"
        "  Dynamic Chunk Request ──► [DelayedFallback (300ms Delay)]\n"
        "    ├── Resolves < 300ms ──► Render component immediately (Zero Flicker)\n"
        "    └── Takes > 300ms    ──► Display animated loading indicator (Smooth UX on Slow 3G)"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BLACK

    # 6. Core Implementation Snippets (Concise as requested)
    add_section_header("6. Core Implementation Snippets")

    p_note = doc.add_paragraph()
    r_n = p_note.add_run("Note: In compliance with faculty guidelines, only the code splitting logic, minimum-delay fallback component, lazy third-party chart integration, and component memoization are documented below.")
    r_n.font.name = FONT_NAME
    r_n.font.size = Pt(10)
    r_n.font.italic = True
    r_n.font.color.rgb = COLOR_BLACK

    snippets = [
        (
            "Snippet 1: Route-Based Lazy Loading with Suspense (src/App.jsx - Steps 3 & 4)",
            "Demonstrates dynamic import() via React.lazy and wrapping Routes inside Suspense:",
            """import { useState, useCallback, lazy, Suspense } from 'react';
import DelayedFallback from './components/DelayedFallback';

// Route-based code splitting
const Home = lazy(() => import('./pages/Home'));
const Projects = lazy(() => import('./pages/Projects'));
const Contact = lazy(() => import('./pages/Contact'));
const NotFound = lazy(() => import('./pages/NotFound'));

function App() {
  const [darkMode, setDarkMode] = useState(false);
  const toggleTheme = useCallback(() => setDarkMode((prev) => !prev), []);

  return (
    <div className={`portfolio-app ${darkMode ? 'dark-theme' : 'light-theme'}`}>
      <NavBar onToggleTheme={toggleTheme} />
      <div className="main-content-wrapper">
        <Suspense fallback={<DelayedFallback delay={300} label="Loading page resources..." />}>
          <Routes>
            <Route path="/" element={<Home />} />
            <Route path="/projects" element={<Projects />} />
            <Route path="/contact" element={<Contact />} />
            <Route path="*" element={<NotFound />} />
          </Routes>
        </Suspense>
      </div>
      <Footer />
    </div>
  );
}"""
        ),
        (
            "Snippet 2: Minimum-Delay Fallback Component (src/components/DelayedFallback.jsx - Supplementary 2)",
            "Demonstrates suppressing the loading indicator for 300ms to eliminate visual flickering on fast networks:",
            """import { useState, useEffect } from 'react';

function DelayedFallback({ delay = 300, label = 'Loading page resources...' }) {
  const [show, setShow] = useState(false);

  useEffect(() => {
    const timer = setTimeout(() => setShow(true), delay);
    return () => clearTimeout(timer);
  }, [delay]);

  // Return null within the initial delay window to prevent visual flickering
  if (!show) return null;

  return (
    <div className="delayed-fallback-wrapper" role="status">
      <div className="lazy-spinner-ring" />
      <p>{label}</p>
    </div>
  );
}
export default DelayedFallback;"""
        ),
        (
            "Snippet 3: On-Demand Lazy-Loaded Chart.js (src/pages/Projects.jsx - Supplementary 1)",
            "Demonstrates deferring heavy library download until explicitly requested by the user:",
            """// Lazy import heavy third-party Chart.js component
const TaskAnalyticsChart = lazy(() => import('../components/TaskAnalyticsChart'));

function Projects() {
  const [showAnalytics, setShowAnalytics] = useState(false);
  
  return (
    <div>
      <button onClick={() => setShowAnalytics((prev) => !prev)}>
        {showAnalytics ? 'Hide Analytics' : '📊 Task Analytics (Lazy Chart)'}
      </button>

      {/* Heavy chunk (205 KB) fetched over network only when showAnalytics is true */}
      {showAnalytics && (
        <Suspense fallback={<DelayedFallback delay={150} label="Loading Chart.js chunk..." />}>
          <TaskAnalyticsChart tasks={tasks} />
        </Suspense>
      )}
    </div>
  );
}"""
        ),
        (
            "Snippet 4: Eliminating Unnecessary Re-renders via React.memo (src/components/Footer.jsx - Supplementary 3)",
            "Demonstrates preventing re-renders of static UI components during parent state transitions:",
            """import React from 'react';

function Footer({ name, year = 2026, email, github }) {
  return (
    <footer className="portfolio-footer">
      <p>© {year} {name}. All rights reserved.</p>
    </footer>
  );
}

// React.memo skips re-rendering when props have not changed
export default React.memo(Footer);"""
        ),
    ]

    for title, desc, code in snippets:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(2)
        r_f = p_title.add_run(title)
        r_f.font.name = FONT_NAME
        r_f.font.size = Pt(11)
        r_f.font.bold = True
        r_f.font.color.rgb = COLOR_BLACK

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(3)
        r_d = p_desc.add_run(desc)
        r_d.font.name = FONT_NAME
        r_d.font.size = Pt(10)
        r_d.font.italic = True
        r_d.font.color.rgb = COLOR_BLACK

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        set_cell_background(c, "F8FAFC")
        c.width = Inches(6.5)
        set_cell_margins(c, top=60, bottom=60, left=100, right=100)

        p_code = c.paragraphs[0]
        p_code.paragraph_format.space_before = Pt(2)
        p_code.paragraph_format.space_after = Pt(2)
        r_c = p_code.add_run(code)
        r_c.font.name = "Consolas"
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = COLOR_BLACK

    # 7. Before vs After Performance Comparison Table
    add_section_header("7. Performance & Bundle Metrics Comparison (Before vs After)")

    p_tbl_intro = doc.add_paragraph()
    r_ti = p_tbl_intro.add_run("Quantitative comparison of JavaScript bundle sizes and network transfer metrics before and after optimization:")
    r_ti.font.name = FONT_NAME
    r_ti.font.size = Pt(10.5)
    r_ti.font.color.rgb = COLOR_BLACK

    test_table = doc.add_table(rows=6, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_headers = ["Metric / Parameter", "Before Optimization", "After Code Splitting", "Improvement & Significance"]
    test_data = [
        ("Initial Entry Bundle Size", "289.91 KB (90.27 KB gzip)", "263.68 KB (84.35 KB gzip)", "Reduced initial bundle; ~26.23 KB core payload reduction upfront."),
        ("Number of JS Bundles", "1 monolithic bundle", "6 modular chunks", "Routes split into dedicated chunks; loaded strictly on demand."),
        ("Home Route Chunk", "Bundled in main", "2.71 KB (0.80 KB gzip)", "Loaded on demand when navigating to '/'."),
        ("Tasks Route Chunk", "Bundled in main", "21.49 KB (6.12 KB gzip)", "Loaded on demand when navigating to '/projects'."),
        ("Heavy Chart.js Chunk", "Bundled upfront if static (~495 KB)", "205.77 KB (70.76 KB gzip)", "Isolated chunk; loaded strictly when analytics button is toggled."),
    ]

    for col_idx, h_text in enumerate(table_headers):
        cell_h = test_table.cell(0, col_idx)
        set_cell_background(cell_h, "F1F5F9")
        p = cell_h.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.name = FONT_NAME
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_BLACK

    for row_idx, row_values in enumerate(test_data):
        for col_idx, val in enumerate(row_values):
            cell_d = test_table.cell(row_idx + 1, col_idx)
            set_cell_background(cell_d, "FFFFFF")
            p = cell_d.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_BLACK

    # React DevTools Profiler Analysis Table
    add_section_header("8. React DevTools Profiler Re-render Analysis")
    p_prof = doc.add_paragraph()
    r_pf = p_prof.add_run("Documentation of unnecessary re-render identified and optimized via React DevTools Profiler (Supplementary Problem 3):")
    r_pf.font.name = FONT_NAME
    r_pf.font.size = Pt(10.5)
    r_pf.font.color.rgb = COLOR_BLACK

    prof_table = doc.add_table(rows=4, cols=4)
    prof_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prof_headers = ["Component", "Triggering Event", "Profiler Finding (Before)", "Optimized Behavior (After)"]
    prof_data = [
        ("Footer.jsx", "Theme toggle (darkMode state in App)", "Re-rendered due to parent render despite unchanged props.", "Skipped render ('Did not render') via React.memo(Footer)."),
        ("NavBar.jsx", "Theme toggle / Route navigation", "Re-created callback caused re-rendering.", "useCallback memoized toggle handler; React.memo skipped render."),
        ("DelayedFallback", "Fast broadband network navigation", "Loading indicator flickers for <50ms.", "Timer suppresses fallback under 300ms; zero visual flicker."),
    ]

    for col_idx, h_text in enumerate(prof_headers):
        cell_h = prof_table.cell(0, col_idx)
        set_cell_background(cell_h, "F1F5F9")
        p = cell_h.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.name = FONT_NAME
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_BLACK

    for row_idx, row_values in enumerate(prof_data):
        for col_idx, val in enumerate(row_values):
            cell_d = prof_table.cell(row_idx + 1, col_idx)
            set_cell_background(cell_d, "FFFFFF")
            p = cell_d.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_BLACK

    # Screenshot placeholder
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(12)
    p_box.paragraph_format.space_after = Pt(8)

    tbl_shot = doc.add_table(rows=1, cols=1)
    tbl_shot.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_shot = tbl_shot.cell(0, 0)
    set_cell_background(c_shot, "F3F4F6")
    c_shot.width = Inches(6.5)
    set_cell_margins(c_shot, top=180, bottom=180, left=150, right=150)
    p_sh = c_shot.paragraphs[0]
    p_sh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sh = p_sh.add_run("[ INSERT SCREENSHOTS OF VITE BUILD CHUNK SIZES, DEVTOOLS NETWORK TAB (SLOW 3G THROTTLING), LAZY CHART.JS MODULE LOADING, & DEVTOOLS PROFILER FLAMEGRAPH HERE ]")
    r_sh.font.name = FONT_NAME
    r_sh.font.size = Pt(11)
    r_sh.font.bold = True
    r_sh.font.color.rgb = COLOR_BLACK

    # 9. Conclusion
    add_section_header("9. Conclusion")
    p_con = doc.add_paragraph()
    r = p_con.add_run(
        "In Practical 8, frontend performance was significantly elevated through route-based code splitting using React.lazy() and Suspense. Heavy third-party libraries (Chart.js, 205 KB) were isolated and loaded on-demand, saving substantial initial bandwidth. A 300ms minimum-delay fallback eliminated loading indicator flickering on high-speed connections, and unnecessary re-renders were diagnosed using the React DevTools Profiler and neutralized via React.memo and useCallback."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    output_path = r"D:\SEM-5\Advance Web Development Frameworks\docs\Practical-08-Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated Practical 8 Word document at: {output_path}")

if __name__ == "__main__":
    create_document()
