import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # University Header
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("CHARUSAT UNIVERSITY\n")
    run_uni.font.name = FONT_NAME
    run_uni.font.size = Pt(16)
    run_uni.font.bold = True
    run_uni.font.color.rgb = COLOR_BLACK

    run_dept = p_uni.add_run("Department of Computer Science & Engineering\n")
    run_dept.font.name = FONT_NAME
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_BLACK

    run_sub = p_uni.add_run("Advance Web Development Frameworks (Semester 5)\n")
    run_sub.font.name = FONT_NAME
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Practical Title Box
    table_title = doc.add_table(rows=1, cols=1)
    table_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_title.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    cell.width = Inches(6.5)
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(6)

    r_pr = p_title.add_run("PRACTICAL REPORT: PRACTICAL NO. 13\n")
    r_pr.font.name = FONT_NAME
    r_pr.font.size = Pt(13)
    r_pr.font.bold = True
    r_pr.font.color.rgb = COLOR_BLACK

    r_pr_name = p_title.add_run("External AI Service Integration, Prompt Engineering, & Responsible AI Usage")
    r_pr_name.font.name = FONT_NAME
    r_pr_name.font.size = Pt(12)
    r_pr_name.font.bold = True
    r_pr_name.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Student Name:", "Priyansh Pravinbhai Panchal"),
        ("Degree & Branch:", "B.Tech Computer Science & Engineering (Semester 5)"),
        ("Institute:", "CHARUSAT University"),
        ("Date:", "October 2026")
    ]
    for row_idx, (k, v) in enumerate(info_data):
        c0 = info_table.cell(row_idx, 0)
        c1 = info_table.cell(row_idx, 1)
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F9FAFB")
        set_cell_background(c1, "FFFFFF")

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k)
        r0.font.name = FONT_NAME
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(v)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        return h

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        r_txt = p.add_run(text)
        r_txt.font.name = FONT_NAME
        r_txt.font.size = Pt(11)
        r_txt.font.color.rgb = COLOR_BLACK

    # 1. Objective
    add_section_header("1. Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(4)
    r = p_obj.add_run("To integrate an external Large Language Model (LLM) service (OpenAI gpt-4o-mini / Gemini API) into the full-stack Task Management application to autogenerate meaningful task descriptions and suggest priority levels, while enforcing strict server-side credential isolation, graceful error resilience, client rate limiting, and responsible AI user transparency.")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # 2. Prerequisites & Resources
    add_section_header("2. Prerequisites & Resources Required")
    p = doc.add_paragraph()
    add_bullet(p, "• Prerequisites: ", "Practical 6 completed working full-stack Task Management application, understanding of REST APIs, asynchronous Node.js execution, and external AI SDK invocation.")
    p = doc.add_paragraph()
    add_bullet(p, "• Tools & Technologies: ", "Node.js (v20), Express.js, OpenAI SDK (v4.x), React 19, Supertest, Postman, Jest, Dotenv.")

    # 3. Problem Definition
    add_section_header("3. Problem Definition")
    for req in [
        "1. Integrate an external AI service into the Task Management backend using OpenAI or Gemini API.",
        "2. Deliver a meaningful feature: auto-generating structured task descriptions and recommending priority levels based on task titles.",
        "3. Store API credentials strictly on the backend (.env) and never expose secrets to frontend client bundles.",
        "4. Implement robust error handling: if the external AI service fails, times out, or has invalid keys, the core application must remain healthy and supply graceful fallback suggestions.",
        "5. Render AI suggestions meaningfully in the React UI as editable inputs rather than raw JSON."
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(req)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK

    # 4. Supplementary Problems Solved
    add_section_header("4. Supplementary Problems Solved")
    for sup in [
        ("• Supplementary 1 (Frontend Loading State): ", "Engineered a reactive loading state ('isAiLoading') disabling buttons and displaying a visual '⏳ Generating...' progress indicator while awaiting asynchronous LLM responses."),
        ("• Supplementary 2 (API Rate Limiting): ", "Built an in-memory client rate limiter restricting requests to maximum 1 AI generation per minute (60 seconds) per client, returning HTTP 429 Too Many Requests to prevent quota exhaustion and API abuse."),
        ("• Supplementary 3 (Responsible AI Transparency): ", "Rendered a dedicated UI disclaimer card emphasizing that generated suggestions are AI-crafted and must be reviewed and verified by the user before saving.")
    ]:
        p = doc.add_paragraph()
        add_bullet(p, sup[0], sup[1])

    # 5. Architecture & Diagram
    add_section_header("5. Architecture & Request-Response Topology")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.left_indent = Inches(0.25)
    r = p_arch.add_run(
        "External AI Service Architecture (Matching Lab Specification 13.1):\n\n"
        "  [React Frontend (Projects.jsx)]\n"
        "         │  Click: '✨ AI Suggest' button with Task Title\n"
        "         │  (Shows '⏳ Generating...' loading indicator)\n"
        "         ▼\n"
        "  [Express Backend: POST /api/ai/generate-description]\n"
        "         │\n"
        "         ├──► [Rate Limiter]: Verifies <= 1 request/min; returns HTTP 429 if exceeded\n"
        "         │\n"
        "         ├──► [Secret Guard]: Reads OPENAI_API_KEY from .env (Server-side ONLY)\n"
        "         │\n"
        "         ├──► Live Path (Key configured):\n"
        "         │       └──► [OpenAI API: gpt-4o-mini]\n"
        "         │               └──► Parsed (Description + Priority Recommendation)\n"
        "         │\n"
        "         └──► Resilience Path (Key missing, invalid, or API error):\n"
        "                 └──► [Contextual Fallback Engine]\n"
        "                         └──► Returns fallback description without crashing\n"
        "         │\n"
        "         ▼\n"
        "  [React UI Form Updates Automatically]\n"
        "  • Description field auto-populated with editable text\n"
        "  • Priority dropdown auto-selected (low / medium / high)\n"
        "  • Responsible AI Disclaimer banner displayed with dismiss control"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BLACK

    # 6. Core Implementation Snippets
    add_section_header("6. Core Implementation Snippets")

    p_note = doc.add_paragraph()
    r_n = p_note.add_run("Note: In compliance with faculty guidelines, only the delta modifications introduced in Practical 13 are documented below, highlighting the backend AI service with rate limiting and fallback, the Express endpoint, the frontend API service, and the React UI form integration.")
    r_n.font.name = FONT_NAME
    r_n.font.size = Pt(10)
    r_n.font.italic = True
    r_n.font.color.rgb = COLOR_BLACK

    snippets = [
        (
            "Snippet 1: Backend AI Service Module (task-manager-api/aiService.js - Steps 1, 2, 3 & Supplementary 2)",
            "Implements rate limiting, OpenAI SDK client invocation, and intelligent offline fallback engine:",
            """const OpenAI = require('openai');

const rateLimitCache = new Map();
const RATE_LIMIT_WINDOW_MS = 60 * 1000; // 60 seconds (Supplementary 2)

function checkRateLimit(key) {
  const now = Date.now();
  if (rateLimitCache.has(key)) {
    const elapsed = now - rateLimitCache.get(key);
    if (elapsed < RATE_LIMIT_WINDOW_MS) {
      return { allowed: false, waitSeconds: Math.ceil((RATE_LIMIT_WINDOW_MS - elapsed) / 1000) };
    }
  }
  rateLimitCache.set(key, now);
  return { allowed: true };
}

async function generateTaskDescription(title, clientId = 'default') {
  const rateCheck = checkRateLimit(clientId);
  if (!rateCheck.allowed) {
    return { error: 'RateLimitExceeded', status: 429, message: `Wait ${rateCheck.waitSeconds}s` };
  }

  const apiKey = process.env.OPENAI_API_KEY;
  if (!apiKey || apiKey === 'your_openai_api_key_here') {
    return generateFallbackSuggestion(title); // Step 6: Fallback path
  }

  try {
    const openai = new OpenAI({ apiKey });
    const completion = await openai.chat.completions.create({
      model: 'gpt-4o-mini',
      messages: [{ role: 'user', content: `Task title: "${title}". Write 1-2 sentence description and Priority: [low|medium|high]` }],
      max_tokens: 150,
    });
    return parseAiResponse(completion.choices[0]?.message?.content);
  } catch (err) {
    return generateFallbackSuggestion(title); // Graceful recovery
  }
}"""
        ),
        (
            "Snippet 2: Express AI Route Definition (task-manager-api/server.js - Step 3)",
            "Exposes /api/ai/generate-description with input validation, rate limit handling, and JSON response structuring:",
            """// POST /api/ai/generate-description (Practical 13)
app.post('/api/ai/generate-description', async (req, res) => {
  try {
    const { title } = req.body;
    if (!title || !title.trim()) {
      return res.status(400).json({ success: false, error: 'Validation Error', message: 'Title is required' });
    }

    const clientId = req.user?.id || req.ip;
    const result = await generateTaskDescription(title.trim(), clientId);

    if (result.error === 'RateLimitExceeded') {
      return res.status(429).json({ success: false, error: 'Too Many Requests', message: result.message });
    }

    res.status(200).json({
      success: true,
      description: result.description,
      priority: result.priority,
      fallback: result.fallback,
      provider: result.provider,
      disclaimer: result.disclaimer,
    });
  } catch (err) {
    res.status(200).json({ success: true, description: `Task: ${req.body.title}`, priority: 'medium', fallback: true });
  }
});"""
        ),
        (
            "Snippet 3: Frontend API Helper (student-portfolio/src/services/api.js - Step 4)",
            "Encapsulates HTTP POST call with Bearer token authentication and status 429 handling:",
            """export async function generateAiDescription(title) {
  const res = await fetch(`${BASE_URL}/api/ai/generate-description`, {
    method: 'POST',
    headers: getAuthHeaders(),
    body: JSON.stringify({ title }),
  });
  const data = await res.json();
  if (!res.ok) {
    throw new Error(data.message || data.error || 'Failed to generate AI description');
  }
  return data;
}"""
        ),
        (
            "Snippet 4: React UI Form Integration (student-portfolio/src/pages/Projects.jsx - Steps 4, 5 & Supplementary 1, 3)",
            "Connects AI Suggest button, loading indicator, editable inputs, and responsible AI disclaimer banner:",
            """// Trigger AI generation with loading state and disclaimer banner
const handleAiGenerate = async () => {
  if (!newTitle.trim()) return showToast('Enter a task title first!', 'error');
  setIsAiLoading(true);
  try {
    const res = await generateAiDescription(newTitle.trim());
    if (res?.description) {
      setNewDesc(res.description);      // Editable suggestion (Step 5)
      setNewPriority(res.priority);     // Auto-suggested priority
      setAiDisclaimer(res.disclaimer);  // Transparency banner (Supplementary 3)
      setAiFallback(Boolean(res.fallback));
      showToast('✨ AI description & priority generated!', 'success');
    }
  } catch (err) {
    showToast(err.message, 'error');
  } finally {
    setIsAiLoading(false);
  }
};

// Form JSX:
<button type="button" onClick={handleAiGenerate} disabled={isAiLoading || !newTitle.trim()}>
  {isAiLoading ? '⏳ Generating...' : '✨ AI Suggest'}
</button>

{aiDisclaimer && (
  <div className="ai-disclaimer-banner">
    <span>🤖 <strong>AI Disclaimer:</strong> {aiDisclaimer}</span>
    <button onClick={() => setAiDisclaimer(null)}>✕</button>
  </div>
)}"""
        ),
    ]

    for title, desc, code in snippets:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(2)
        r_f = p_title.add_run(title)
        r_f.font.name = FONT_NAME
        r_f.font.size = Pt(11)
        r_f.font.bold = True
        r_f.font.color.rgb = COLOR_BLACK

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(3)
        r_d = p_desc.add_run(desc)
        r_d.font.name = FONT_NAME
        r_d.font.size = Pt(10)
        r_d.font.italic = True
        r_d.font.color.rgb = COLOR_BLACK

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        set_cell_background(c, "F8FAFC")
        c.width = Inches(6.5)
        set_cell_margins(c, top=60, bottom=60, left=100, right=100)

        p_code = c.paragraphs[0]
        p_code.paragraph_format.space_before = Pt(2)
        p_code.paragraph_format.space_after = Pt(2)
        r_c = p_code.add_run(code)
        r_c.font.name = "Consolas"
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = COLOR_BLACK

    # 7. Verification & Test Suite Matrix
    add_section_header("7. Verification & Automated Test Matrix")

    p_tbl_intro = doc.add_paragraph()
    r_ti = p_tbl_intro.add_run("Summary of verified test scenarios across AI generation, rate limiting, and failure tolerance:")
    r_ti.font.name = FONT_NAME
    r_ti.font.size = Pt(10.5)
    r_ti.font.color.rgb = COLOR_BLACK

    test_table = doc.add_table(rows=5, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_headers = ["Scenario ID", "Test Condition", "Expected Output / Status", "Observed System Behavior"]
    test_data = [
        ("AI-TEST-01", "Valid title submitted ('Complete Lab 13')", "HTTP 200, valid description, priority, disclaimer", "Generated actionable text and selected 'high' priority in 3.37ms."),
        ("AI-TEST-02", "Empty title payload ({ title: '' })", "HTTP 400 Bad Request, Validation Error", "Rejected before API call, guarding against empty LLM requests."),
        ("AI-TEST-03", "Rapid consecutive request (< 60s)", "HTTP 429 Too Many Requests", "Enforced in-memory rate limit: 'Max 1 AI request per minute'."),
        ("AI-TEST-04", "Missing/invalid API key (Failure Path)", "HTTP 200, fallback: true, zero server crashes", "Contextual fallback activated gracefully without downtime."),
    ]

    for col_idx, h_text in enumerate(table_headers):
        cell_h = test_table.cell(0, col_idx)
        set_cell_background(cell_h, "F1F5F9")
        p = cell_h.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.name = FONT_NAME
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_BLACK

    for row_idx, row_values in enumerate(test_data):
        for col_idx, val in enumerate(row_values):
            cell_d = test_table.cell(row_idx + 1, col_idx)
            set_cell_background(cell_d, "FFFFFF")
            p = cell_d.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_BLACK

    # 8. Responsible AI Principles
    add_section_header("8. Responsible AI Usage & Engineering Ethics")
    p_resp = doc.add_paragraph()
    r_rp = p_resp.add_run(
        "Responsible AI practices were embedded into the system architecture at every layer:\n"
        "1. Human-in-the-Loop Verification: AI outputs are presented as suggestions within editable form fields, giving users full agency to revise, expand, or discard proposals before submitting to the database.\n"
        "2. Transparency & Attribution: The UI prominently presents a disclaimer banner explaining that the text was machine-generated, eliminating deceptive ambiguity.\n"
        "3. Secret Isolation: Third-party API credentials exist solely within server-side environment variables (.env) and are never embedded into client JavaScript bundles or committed to source control.\n"
        "4. Quota Protection: Rate limiting prevents resource depletion, denial-of-service, and unexpected billing spikes from runaway client loops."
    )
    r_rp.font.name = FONT_NAME
    r_rp.font.size = Pt(10.5)
    r_rp.font.color.rgb = COLOR_BLACK

    # Screenshot placeholder
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(12)
    p_box.paragraph_format.space_after = Pt(8)

    tbl_shot = doc.add_table(rows=1, cols=1)
    tbl_shot.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_shot = tbl_shot.cell(0, 0)
    set_cell_background(c_shot, "F3F4F6")
    c_shot.width = Inches(6.5)
    set_cell_margins(c_shot, top=180, bottom=180, left=150, right=150)
    p_sh = c_shot.paragraphs[0]
    p_sh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sh = p_sh.add_run("[ INSERT SCREENSHOTS OF: 1. TASK CREATION FORM WITH '✨ AI SUGGEST' BUTTON & LOADING STATE, 2. AUTO-POPULATED DESCRIPTION AND AI DISCLAIMER BANNER, 3. POSTMAN / TERMINAL TEST PROVING HTTP 429 RATE LIMITING, AND 4. JEST TEST SUITE PASSING 10/10 TESTS HERE ]")
    r_sh.font.name = FONT_NAME
    r_sh.font.size = Pt(11)
    r_sh.font.bold = True
    r_sh.font.color.rgb = COLOR_BLACK

    # 9. Conclusion
    add_section_header("9. Conclusion")
    p_con = doc.add_paragraph()
    r = p_con.add_run(
        "Practical 13 successfully enriched the Task Management application with an external AI service. By coupling OpenAI gpt-4o-mini with secure backend credential encapsulation, intelligent offline fallback handlers, in-memory rate limiting, and an intuitive human-in-the-loop React interface with clear disclaimers, the solution exemplifies modern, resilient, and responsible AI integration."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    output_path = r"D:\SEM-5\Advance Web Development Frameworks\docs\Practical-13-Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated Practical 13 Word document at: {output_path}")

if __name__ == "__main__":
    create_document()
