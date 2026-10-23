import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # University Header
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_uni = p_uni.add_run("CHARUSAT UNIVERSITY\n")
    run_uni.font.name = FONT_NAME
    run_uni.font.size = Pt(16)
    run_uni.font.bold = True
    run_uni.font.color.rgb = COLOR_BLACK

    run_dept = p_uni.add_run("Department of Computer Science & Engineering\n")
    run_dept.font.name = FONT_NAME
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_BLACK

    run_sub = p_uni.add_run("Advance Web Development Frameworks (Semester 5)\n")
    run_sub.font.name = FONT_NAME
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Practical Title Box
    table_title = doc.add_table(rows=1, cols=1)
    table_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_title.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    cell.width = Inches(6.5)
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(6)

    r_pr = p_title.add_run("PRACTICAL REPORT: PRACTICAL NO. 4\n")
    r_pr.font.name = FONT_NAME
    r_pr.font.size = Pt(13)
    r_pr.font.bold = True
    r_pr.font.color.rgb = COLOR_BLACK

    r_pr_name = p_title.add_run("RESTful Backend Server with Express Middleware Pipeline and In-Memory CRUD")
    r_pr_name.font.name = FONT_NAME
    r_pr_name.font.size = Pt(12)
    r_pr_name.font.bold = True
    r_pr_name.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Student Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Student Name:", "Priyansh Pravinbhai Panchal"),
        ("Degree & Branch:", "B.Tech Computer Science & Engineering (Semester 5)"),
        ("Institute:", "CHARUSAT University"),
        ("Date:", "October 2026")
    ]
    for row_idx, (k, v) in enumerate(info_data):
        c0 = info_table.cell(row_idx, 0)
        c1 = info_table.cell(row_idx, 1)
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F9FAFB")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k)
        r0.font.name = FONT_NAME
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(v)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        return h

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        r_txt = p.add_run(text)
        r_txt.font.name = FONT_NAME
        r_txt.font.size = Pt(11)
        r_txt.font.color.rgb = COLOR_BLACK

    # 1. Objective
    add_section_header("1. Objective")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(4)
    r = p_obj.add_run("To design and implement a RESTful backend server with complete CRUD endpoints using an Express.js middleware pipeline, in-memory array storage, and structured error handling.")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    # 2. Prerequisites & Resources
    add_section_header("2. Prerequisites & Resources Required")
    p = doc.add_paragraph()
    add_bullet(p, "• Prerequisites: ", "Basic JavaScript (functions, objects, arrays, async basics), HTTP methods (GET, POST, PUT, DELETE) and status code semantics.")
    p = doc.add_paragraph()
    add_bullet(p, "• Tools & Technologies: ", "Node.js (v24.x), npm (v11.x), Express.js (v4.21.x), Visual Studio Code, Postman / Thunder Client / cURL.")

    # 3. Problem Definition
    add_section_header("3. Problem Definition")
    for req in [
        "1. Build a Node/Express backend for a Task Management system (in 'task-manager-api').",
        "2. Implement REST endpoints for creating, reading, updating, and deleting tasks using an in-memory array.",
        "3. Apply a request logging middleware that logs method, URL, and timestamp for every incoming request.",
        "4. Apply a global error handling middleware as the last middleware in the pipeline.",
        "5. Use correct HTTP status codes (200 OK, 201 Created, 400 Bad Request, 404 Not Found, 415 Unsupported Media Type, 500 Internal Server Error) for all responses."
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(req)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK

    # 4. Supplementary Problems Solved
    add_section_header("4. Supplementary Problems Solved")
    for sup in [
        ("• Supplementary 1: ", "Added middleware that inspects POST/PUT requests and rejects those missing the 'Content-Type: application/json' header with HTTP 415 Unsupported Media Type."),
        ("• Supplementary 2: ", "Added route-specific middleware (validateTaskId) that validates task ID format (positive integer) before reaching controller functions, returning HTTP 400 Bad Request if invalid."),
        ("• Supplementary 3: ", "Added a 404 catch-all handler for undefined routes that returns a structured JSON response instead of default HTML.")
    ]:
        p = doc.add_paragraph()
        add_bullet(p, sup[0], sup[1])

    # 5. Architecture & Middleware Pipeline
    add_section_header("5. Architecture & Middleware Pipeline")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.left_indent = Inches(0.25)
    r = p_arch.add_run(
        "Client (Postman / Browser / cURL)\n"
        "  │\n"
        "  ▼ [1. Global Request Logging Middleware: Method + URL + ISO Timestamp]\n"
        "  │\n"
        "  ▼ [2. Content-Type Header Validator Middleware (POST/PUT check)]\n"
        "  │\n"
        "  ▼ [3. express.json() Body Parser Middleware]\n"
        "  │\n"
        "  ▼ Express REST Router (CRUD Endpoints)\n"
        "  │   ├── GET    /tasks         → Return all tasks (200 OK)\n"
        "  │   ├── GET    /tasks/:id     → [validateTaskId] → Return task by ID (200 or 404)\n"
        "  │   ├── POST   /tasks         → Create new task (201 Created)\n"
        "  │   ├── PUT    /tasks/:id     → [validateTaskId] → Update task by ID (200 or 404)\n"
        "  │   └── DELETE /tasks/:id     → [validateTaskId] → Remove task by ID (200 or 404)\n"
        "  │\n"
        "  ▼ [4. 404 Catch-All Handler: Structured JSON response for undefined routes]\n"
        "  │\n"
        "  ▼ [5. Global Error Handling Middleware: (err, req, res, next) → 500 JSON response]"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BLACK

    # 6. Core Implementation Snippets (Concise as requested)
    add_section_header("6. Core Implementation Snippets")

    p_note = doc.add_paragraph()
    r_n = p_note.add_run("Note: Only key middleware definitions and core CRUD logic from 'task-manager-api/server.js' are documented below to maintain a concise submission report.")
    r_n.font.name = FONT_NAME
    r_n.font.size = Pt(10)
    r_n.font.italic = True
    r_n.font.color.rgb = COLOR_BLACK

    snippets = [
        (
            "Snippet 1: Global Request Logging Middleware (server.js)",
            "Demonstrates intercepting every incoming HTTP request and logging method, URL, and timestamp:",
            """// Step 3: Global Request Logging Middleware
app.use((req, res, next) => {
  const timestamp = new Date().toISOString();
  console.log(`[${timestamp}] ${req.method} ${req.url}`);
  next();
});"""
        ),
        (
            "Snippet 2: Content-Type Validation Middleware (Supplementary Problem 1)",
            "Demonstrates enforcing Content-Type: application/json on state-modifying requests (POST/PUT):",
            """// Supplementary Problem 1: Header Validator Middleware
const requireJsonContent = (req, res, next) => {
  if (['POST', 'PUT', 'PATCH'].includes(req.method)) {
    const contentType = req.headers['content-type'];
    if (!contentType || !contentType.toLowerCase().includes('application/json')) {
      return res.status(415).json({
        success: false,
        error: 'Unsupported Media Type',
        message: 'Requests with payload (POST/PUT) must include Content-Type: application/json header'
      });
    }
  }
  next();
};
app.use(requireJsonContent);"""
        ),
        (
            "Snippet 3: Route-Specific ID Validation Middleware (Supplementary Problem 2)",
            "Demonstrates validating URL parameter format before executing the controller logic:",
            """// Supplementary Problem 2: Route-Specific Task ID Validator
const validateTaskId = (req, res, next) => {
  const { id } = req.params;
  const parsedId = parseInt(id, 10);

  if (isNaN(parsedId) || parsedId <= 0 || !/^\\d+$/.test(id)) {
    return res.status(400).json({
      success: false,
      error: 'Invalid Task ID',
      message: `Task ID '${id}' must be a valid positive integer.`
    });
  }
  req.taskId = parsedId;
  next();
};"""
        ),
        (
            "Snippet 4: RESTful CRUD Controller Endpoints (server.js)",
            "Demonstrates creating, reading, updating, and deleting tasks with in-memory storage and status codes:",
            """// GET /tasks - Read all tasks (200 OK)
app.get('/tasks', (req, res) => {
  res.status(200).json({ success: true, count: tasks.length, data: tasks });
});

// POST /tasks - Create task (201 Created)
app.post('/tasks', (req, res) => {
  const { title, description = '', completed = false } = req.body;
  if (!title || !title.trim()) {
    return res.status(400).json({ success: false, error: 'Task title is required.' });
  }
  const newTask = { id: nextId++, title: title.trim(), description: description.trim(), completed: Boolean(completed), createdAt: new Date().toISOString() };
  tasks.push(newTask);
  res.status(201).json({ success: true, message: 'Task created successfully', data: newTask });
});

// PUT /tasks/:id - Update task (200 OK or 404)
app.put('/tasks/:id', validateTaskId, (req, res) => {
  const taskIndex = tasks.findIndex((t) => t.id === req.taskId);
  if (taskIndex === -1) return res.status(404).json({ success: false, error: 'Task Not Found' });
  const { title, description, completed } = req.body;
  tasks[taskIndex] = { ...tasks[taskIndex], ...req.body, updatedAt: new Date().toISOString() };
  res.status(200).json({ success: true, message: 'Task updated successfully', data: tasks[taskIndex] });
});

// DELETE /tasks/:id - Delete task (200 OK or 404)
app.delete('/tasks/:id', validateTaskId, (req, res) => {
  const taskIndex = tasks.findIndex((t) => t.id === req.taskId);
  if (taskIndex === -1) return res.status(404).json({ success: false, error: 'Task Not Found' });
  const deletedTask = tasks.splice(taskIndex, 1)[0];
  res.status(200).json({ success: true, message: 'Task deleted successfully', data: deletedTask });
});"""
        ),
        (
            "Snippet 5: 404 Handler & Global Error Middleware (Step 5 & Supplementary 3)",
            "Demonstrates handling undefined routes with JSON and catching internal errors at pipeline end:",
            """// Supplementary Problem 3: 404 Handler for Undefined Routes
app.use((req, res) => {
  res.status(404).json({
    success: false,
    error: 'Route Not Found',
    message: `The requested endpoint '${req.method} ${req.originalUrl}' does not exist on this server.`
  });
});

// Step 5: Global Error Handling Middleware (4 arguments)
app.use((err, req, res, next) => {
  console.error('[Global Error Handler]:', err.message);
  res.status(500).json({
    success: false,
    error: 'Internal Server Error',
    message: err.message || 'Something went wrong on the server'
  });
});"""
        ),
    ]

    for title, desc, code in snippets:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(8)
        p_title.paragraph_format.space_after = Pt(2)
        r_f = p_title.add_run(title)
        r_f.font.name = FONT_NAME
        r_f.font.size = Pt(11)
        r_f.font.bold = True
        r_f.font.color.rgb = COLOR_BLACK

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(3)
        r_d = p_desc.add_run(desc)
        r_d.font.name = FONT_NAME
        r_d.font.size = Pt(10)
        r_d.font.italic = True
        r_d.font.color.rgb = COLOR_BLACK

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        set_cell_background(c, "F8FAFC")
        c.width = Inches(6.5)
        set_cell_margins(c, top=60, bottom=60, left=100, right=100)

        p_code = c.paragraphs[0]
        p_code.paragraph_format.space_before = Pt(2)
        p_code.paragraph_format.space_after = Pt(2)
        r_c = p_code.add_run(code)
        r_c.font.name = "Consolas"
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = COLOR_BLACK

    # 7. Output & Verification Table
    add_section_header("7. Output & Verification Results")
    
    p_tbl_intro = doc.add_paragraph()
    r_ti = p_tbl_intro.add_run("Summary of tested endpoints, middleware interception, and status code verification:")
    r_ti.font.name = FONT_NAME
    r_ti.font.size = Pt(10.5)
    r_ti.font.color.rgb = COLOR_BLACK

    test_table = doc.add_table(rows=8, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_headers = ["Endpoint", "Method", "Status Code", "Observed Response / Behavior"]
    test_data = [
        ("GET /tasks", "GET", "200 OK", "Returns all 4 in-memory tasks with count."),
        ("POST /tasks", "POST", "201 Created", "Creates task #5 and returns the created JSON object."),
        ("POST /tasks (No Header)", "POST", "415 Unsupported", "Rejected by requireJsonContent middleware without JSON header."),
        ("GET /tasks/:id (Invalid)", "GET", "400 Bad Request", "Rejected by validateTaskId for invalid parameter ('invalid-id')."),
        ("PUT /tasks/5", "PUT", "200 OK", "Updates task #5 'completed' to true with updatedAt timestamp."),
        ("DELETE /tasks/5", "DELETE", "200 OK", "Removes task #5; subsequent call returns 404 Not Found."),
        ("GET /api/undefined", "GET", "404 Not Found", "Catches unhandled route and returns structured JSON."),
    ]

    for col_idx, h_text in enumerate(table_headers):
        cell_h = test_table.cell(0, col_idx)
        set_cell_background(cell_h, "F1F5F9")
        p = cell_h.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.name = FONT_NAME
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_BLACK

    for row_idx, row_values in enumerate(test_data):
        for col_idx, val in enumerate(row_values):
            cell_d = test_table.cell(row_idx + 1, col_idx)
            set_cell_background(cell_d, "FFFFFF")
            p = cell_d.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_BLACK

    # Screenshot placeholder
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(12)
    p_box.paragraph_format.space_after = Pt(8)
    
    tbl_shot = doc.add_table(rows=1, cols=1)
    tbl_shot.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_shot = tbl_shot.cell(0, 0)
    set_cell_background(c_shot, "F3F4F6")
    c_shot.width = Inches(6.5)
    set_cell_margins(c_shot, top=180, bottom=180, left=150, right=150)
    p_sh = c_shot.paragraphs[0]
    p_sh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sh = p_sh.add_run("[ INSERT POSTMAN / THUNDER CLIENT SCREENSHOTS OF CRUD REQUESTS & MIDDLEWARE LOGS HERE ]")
    r_sh.font.name = FONT_NAME
    r_sh.font.size = Pt(11)
    r_sh.font.bold = True
    r_sh.font.color.rgb = COLOR_BLACK

    # 8. Conclusion
    add_section_header("8. Conclusion")
    p_con = doc.add_paragraph()
    r = p_con.add_run(
        "In Practical 4, an Express.js backend server was constructed to implement a RESTful Task Management API. An end-to-end middleware pipeline was built including global request logging, Content-Type enforcement, route-specific ID validation, in-memory CRUD operations, structured 404 routing, and global error handling. All HTTP status codes were validated according to REST best practices."
    )
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK

    output_path = r"D:\SEM-5\Advance Web Development Frameworks\docs\Practical-04-Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated Practical 4 Word document at: {output_path}")

if __name__ == "__main__":
    create_document()
